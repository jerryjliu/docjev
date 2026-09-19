"""Generate the CC0 synthetic corpus. No model, network, or credentials are used.

Run with the datasets extra installed. Committed PowerPoint sources are authored
separately; this generator verifies their declared hashes rather than editing them.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import random
import tempfile
import zipfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageFilter
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import HexColor, white
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
SEED = 20260918
VERSION = "1.0.0"
CLASS_LABELS = ["invoice", "purchase_order", "contract", "resume", "pitch_deck", "other"]
SPLIT_LABELS = ["claim_form", "incident_report", "repair_estimate", "invoice", "correspondence", "other"]
PURPLE = "#3E18F9"
INK = "#111111"
GRAY = "#737373"
PALETTE = ["#3E18F9", "#4B72FE", "#000000", "#3E18F9"]
ORG = ["Cedar Bay", "Juniper Works", "Morrow Studio", "Holloway Labs", "Cobalt Harbor", "Quarry Lane", "Willow Ridge", "Marble Kite", "North Fern", "Orchid Field", "Silver Loom", "Maple Current"]
PEOPLE = ["Avery Vale", "Jordan Lark", "Casey Rowan", "Morgan Wren", "Taylor Reed", "Riley Finch", "Skyler Brooks", "Parker Quinn", "Hayden Lake", "Drew Linden"]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fonts() -> None:
    for name, file in [("Overused", "OverusedGrotesk-Regular.ttf"), ("OverusedMedium", "OverusedGrotesk-Medium.ttf"), ("Plex", "IBMPlexMono-Regular.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(ROOT / "datasets/assets/fonts" / file)))


def neutral_zip(path: Path) -> None:
    """Fix container timestamps so source file bytes reproduce across runs."""
    source = zipfile.ZipFile(path)
    parts = {name: source.read(name) for name in source.namelist()}
    source.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in sorted(parts.items()):
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)


def paragraph(c: canvas.Canvas, text: str, x: float, y: float, width: float, *, size: float = 11, color: str = INK, bold: bool = False) -> float:
    style = ParagraphStyle("body", fontName="OverusedMedium" if bold else "Overused", fontSize=size, leading=size * 1.48, textColor=HexColor(color), spaceAfter=8)
    p = Paragraph(escape(text), style)
    _, height = p.wrap(width, 700)
    p.drawOn(c, x, y - height)
    return y - height - 12


def content(category: str, index: int, *, pages: int = 1, domain: str = "classification") -> list[dict]:
    """Human-authored content templates; labels follow source identity, not models."""
    org = ORG[index % len(ORG)]
    client = ORG[(index + 5) % len(ORG)]
    person = PEOPLE[index % len(PEOPLE)]
    ref = f"{7300 + index}"
    value = 620 + (index % 17) * 83
    day = index % 25 + 1
    date = f"September {day}, 2026"
    base = {"organization": f"{org} (fictional)", "reference": ref, "date": date}
    templates = {
        "invoice": {"title": "Invoice", "subtitle": f"{org} Services to {client}", "fields": [("Invoice number", f"INV-{ref}"), ("Bill to", f"{client}, 14 Example Way"), ("Issued", date), ("Payment terms", "Net 30, USD")], "sections": [("Payment requested", f"Please remit ${value:,.2f} for the completed services listed below. Reference INV-{ref} with your payment. No account or payment details in this fixture are real."), ("Remittance note", "Payment is due thirty days after issue. Contact billing@example.invalid for a corrected copy.")], "table": [["Service", "Quantity", "Amount"], ["Equipment restoration", "1", f"${value-180:,.2f}"], ["Inspection and travel", "1", "$180.00"], ["TOTAL DUE", "", f"${value:,.2f}"]]},
        "purchase_order": {"title": "Purchase order", "subtitle": f"{org} Procurement", "fields": [("Order number", f"PO-{ref}"), ("Supplier", client), ("Requested delivery", "October 15, 2026"), ("Ship to", "Receiving desk, 18 Example Lane")], "sections": [("Authorization to supply", "The buyer authorizes the supplier to deliver the items below under the agreed commercial terms. Send an order acknowledgement before dispatch."), ("Receiving instructions", f"Reference PO-{ref} on the packing slip. Deliver between 09:00 and 16:00. The purchasing contact is {person}.")], "table": [["Item", "Quantity", "Unit price"], ["Sensor mounting kit", "12", "$42.00"], ["Weatherproof enclosure", "6", "$95.00"], ["Estimated order total", "", "$1,074.00"]]},
        "contract": {"title": "Service agreement", "subtitle": f"{org} and {client}", "fields": [("Agreement reference", f"AGR-{ref}"), ("Effective date", date), ("Customer", org), ("Provider", client)], "sections": [("Scope and delivery", "The provider agrees to inspect and maintain the customer's demonstration equipment. The customer will provide reasonable access to the equipment and appoint a contact for acceptance."), ("Fees and term", f"The customer agrees to pay a monthly fee of ${value:,.2f}. This agreement runs for twelve months unless either party terminates it on thirty days written notice."), ("Confidentiality and signatures", f"Each party will protect confidential information disclosed during the engagement. Authorized signatories: {person} for the customer and Alex Rowan for the provider. Signature: SAMPLE ONLY.")]},
        "resume": {"title": person, "subtitle": "Operations analyst", "fields": [("Contact", f"candidate{index}@example.invalid"), ("Location", "Fictional City"), ("Portfolio", "portfolio.example.invalid"), ("Availability", "November 2026")], "sections": [("Professional summary", "Operations analyst with experience documenting service processes and maintaining reporting systems for small manufacturing teams. Seeking a role improving the quality of operational data."), ("Experience", f"{org}, Operations analyst, 2023-2026. Built weekly service reports and reduced missing work-order fields. {client}, Project coordinator, 2020-2023. Scheduled equipment inspections and tracked supplier delivery commitments."), ("Education and skills", "Bachelor of Science in Industrial Engineering, Example Institute, 2020. Skills: SQL, spreadsheets, process mapping, stakeholder interviews, and technical writing.")]},
        "pitch_deck": {"title": f"{org} investor overview", "subtitle": "Seed financing presentation", "fields": [("Product", "Equipment monitoring software"), ("Audience", "Prospective investors"), ("Stage", "Synthetic seed company"), ("Date", date)], "sections": [("Customer problem", "Independent service teams track equipment condition across fragmented spreadsheets. Missing maintenance history causes unnecessary site visits and delays repairs."), ("Product and business model", "Our subscription software combines equipment logs with a shared maintenance calendar. Illustrative annual contract price: $12,000 per team. All customer and financial figures in this document are fictional."), ("Financing request", "Seeking an illustrative $2 million seed round to fund product development and a twelve-month customer pilot. The founding team combines field operations and software engineering experience.")]},
        "other": {"title": "Community observatory field guide", "subtitle": f"{org} Learning Circle", "fields": [("Edition", "Autumn 2026"), ("Topic", "Night sky observations"), ("Prepared by", person), ("Reference", f"NOTE-{ref}")], "sections": [("Before the session", "Choose an open observation site away from bright street lights. Give your eyes time to adjust before looking for faint stars. A red light helps preserve night vision."), ("Recording observations", "Write the time, sky conditions, and direction of each observation. Sketch the position of bright objects against a familiar constellation. Compare notes at the end of the session."), ("Equipment care", "Keep protective covers on unused lenses. Dry the exterior before storage and leave damp cases open until the lining is dry.")]},
        "claim_form": {"title": "Property damage claim", "subtitle": f"{org} Mutual demonstration form", "fields": [("Claim reference", f"CL-{ref}"), ("Policyholder", person), ("Loss date", date), ("Location", "21 Fictional Lane")], "sections": [("Statement of loss", "A water supply fitting failed in the service room during the evening. Water damaged the floor finish and a storage cabinet. The policyholder isolated the water supply and notified the building manager."), ("Requested reimbursement", f"The policyholder requests reimbursement of documented repair expenses. Estimated loss: ${value * 4:,.2f}. Supporting reports and vendor records are attached separately."), ("Declaration", f"I certify that this fictional statement represents the circumstances described in this demonstration. Policyholder: {person}. Signature: SAMPLE ONLY.")]},
        "incident_report": {"title": "Incident report", "subtitle": f"{org} Facilities team", "fields": [("Report number", f"IR-{ref}"), ("Incident date", date), ("Reporting officer", person), ("Site", "21 Fictional Lane, service room")], "sections": [("Observed conditions", "At 18:20 the facilities officer found standing water next to the supply cabinet. The floor covering was wet along the north wall. No injuries were reported."), ("Immediate response", "The officer closed the isolation valve, placed warning signs at the entrance, and moved stored materials away from the wet area. A repair contractor attended the following morning."), ("Cause and follow up", "A split flexible supply hose was visible behind the cabinet. Retain the damaged fitting for inspection and record moisture readings before reinstating the floor.")]},
        "repair_estimate": {"title": "Repair estimate", "subtitle": f"{org} Restoration", "fields": [("Estimate reference", f"EST-{ref}"), ("Prepared for", person), ("Site", "21 Fictional Lane"), ("Valid through", "October 30, 2026")], "sections": [("Proposed work", "Remove the damaged floor finish, dry the substrate, and replace the affected panels. Reinstate cabinet skirting after moisture readings meet the installation requirement."), ("Conditions", "This quotation is an estimate for proposed work. It is subject to customer acceptance before scheduling. Additional concealed damage requires a revised written estimate.")], "table": [["Proposed work", "Hours", "Estimated cost"], ["Removal and drying", "8", "$960.00"], ["Floor finish and skirting", "12", "$1,440.00"], ["Estimated total", "", "$2,400.00"]]},
        "correspondence": {"title": "Claim correspondence", "subtitle": f"{org} Claims desk", "fields": [("To", person), ("From", "claims@example.invalid"), ("Date", date), ("Subject", f"Supporting information for CL-{ref}")], "sections": [("Dear policyholder", "Thank you for providing the supporting documents. We have received the initial loss statement and contractor records. The assigned reviewer will compare the records with the policy terms."), ("Next step", "Please retain the damaged materials until the inspection is complete. If you locate additional photographs from the day of the loss, send copies with your claim reference."), ("Regards", "Alex Rowan, demonstration claims coordinator. This letter contains no coverage decision and belongs to a synthetic example.")]},
    }
    first = {**base, **templates[category]}
    out = [first]
    for _page in range(2, pages + 1):
        if category == "invoice":
            sections = [("Service detail", f"Continuation of INV-{ref}. Technician work log: isolate damaged fittings, complete pressure testing, and confirm the equipment is safe to return to service."), ("Work completed", "Inspection: 1.5 hours. Restoration: 3.0 hours. Verification: 1.0 hour. The charges summarize work already delivered and form part of the amount due on the preceding page.")]
        elif category == "contract":
            sections = [("Liability and termination", f"Continuation of agreement AGR-{ref}. Each party remains responsible for its own acts. Either party may terminate for a material breach that remains unresolved thirty days after written notice."), ("Governing terms", "Changes require a written amendment signed by both parties. This fictional agreement creates no obligation and is provided only as a document processing example.")]
        elif category == "incident_report":
            sections = [("Follow up observations", f"Continuation of IR-{ref}. The second inspection found no further escape of water. A moisture meter showed the greatest reading near the cabinet base."), ("Witness statement", f"{person} reported hearing water flow before entering the room. The witness contacted the facilities desk immediately. Photographic evidence is represented by the written site notes in this synthetic record.")]
        elif category == "repair_estimate":
            sections = [("Work schedule and exclusions", f"Continuation of estimate EST-{ref}. Allow two days for drying and one day for reinstatement. The quoted scope excludes structural repair and work outside the affected room."), ("Acceptance", "The customer may approve the proposed scope in writing. Materials will be ordered after acceptance. This estimate is valid until the date stated on its first page.")]
        elif category == "claim_form":
            sections = [("Additional loss details", f"Continuation of CL-{ref}. The affected cabinet contained spare equipment and cleaning materials. None of the items are personal records or genuine customer property."), ("Supporting statement", "The policyholder had not observed a leak before the incident. An equipment inspection two weeks earlier recorded normal operation. Repairs await the contractor's final scope.")]
        elif category == "pitch_deck":
            sections = [("Go to market", "Start with independent service firms managing between fifty and five hundred assets. Run a small paid pilot with a named operations lead and review product usage each month."), ("Illustrative financial plan", "Year one: $120,000 annual recurring revenue. Year two: $420,000. Year three: $960,000. These are invented planning figures, not operating results or forecasts for a real company.")]
        else:
            sections = [("Additional notes", f"Continuation of reference {ref}. These supporting details belong to the same source document as the preceding page."), ("Record keeping", "Maintain the original reference with all supporting notes. This page intentionally omits the first-page title to exercise continuity across a page boundary.")]
        out.append({**base, "title": "", "subtitle": f"Reference {ref} / continued", "fields": [], "sections": sections})
    return out


def draw_page(c: canvas.Canvas, page: dict, family: str, number: int, total: int) -> None:
    width, height = 612, 792
    family_number = int(family.rsplit("-", 1)[-1])
    family_id = family_number % 2 if "-dev-" in family else family_number % 4 + 2
    accent = PALETTE[family_id % 4]
    margin = {0: 52, 1: 56, 2: 66, 3: 52, 4: 44, 5: 60}[family_id]
    usable = width - 2 * margin
    c.setFillColor(HexColor("#F5F5F5"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    if family_id == 0:
        c.setFillColor(HexColor(accent))
        c.rect(0, height - 11, width, 11, fill=1, stroke=0)
    elif family_id == 1:
        c.setFillColor(HexColor(accent))
        c.rect(0, 0, 13, height, fill=1, stroke=0)
    elif family_id == 3:
        c.setFillColor(HexColor("#E9E5FF"))
        c.rect(0, height - 174, width, 174, fill=1, stroke=0)
    elif family_id == 4:
        c.setStrokeColor(HexColor(accent))
        c.setLineWidth(2)
        c.line(margin, height - 57, width - margin, height - 57)
    elif family_id == 5:
        c.setFillColor(HexColor("#37D7FA"))
        c.rect(width - 24, 0, 24, height, fill=1, stroke=0)
    c.setFillColor(HexColor(accent))
    c.setFont("Plex", 8)
    c.drawString(margin, height - 44, page["organization"].upper())
    y = height - 76
    if page["title"]:
        y = paragraph(c, page["title"], margin, y, usable, size=30, bold=True)
    y = paragraph(c, page["subtitle"], margin, y, usable, size=12, color=GRAY)
    y -= 16
    fields = page.get("fields", [])
    if fields:
        for i in range(0, len(fields), 2):
            for j, (key, val) in enumerate(fields[i:i+2]):
                x = margin + j * (usable / 2 + 6)
                paragraph(c, key.upper(), x, y, usable / 2 - 10, size=8, color=GRAY)
                paragraph(c, val, x, y - 17, usable / 2 - 10, size=11, bold=True)
            y -= 54
        y -= 5
    for heading, body in page.get("sections", []):
        y = paragraph(c, heading, margin, y, usable, size=13, bold=True)
        y = paragraph(c, body, margin, y + 5, usable, size=10.5)
        y -= 5
    if table := page.get("table"):
        tab = Table(table, colWidths=[usable * .59, usable * .16, usable * .25])
        tab.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "Overused"), ("FONTNAME", (0, 0), (-1, 0), "OverusedMedium"), ("FONTSIZE", (0, 0), (-1, -1), 10), ("BACKGROUND", (0, 0), (-1, 0), HexColor(accent)), ("TEXTCOLOR", (0, 0), (-1, 0), white), ("LINEBELOW", (0, 1), (-1, -1), .5, HexColor("#D6D6D6")), ("TOPPADDING", (0, 0), (-1, -1), 11), ("BOTTOMPADDING", (0, 0), (-1, -1), 11)]))
        _, h = tab.wrap(usable, 500)
        tab.drawOn(c, margin, y - h)
        y -= h
    if y < 80:
        raise ValueError(f"Content overflow in {page['reference']}: y={y}")
    c.setStrokeColor(HexColor("#D2D2D2"))
    c.line(margin, 52, width - margin, 52)
    c.setFillColor(HexColor(GRAY))
    c.setFont("Plex", 7)
    c.drawString(margin, 35, "SYNTHETIC EXAMPLE / NO REAL PERSON OR TRANSACTION")
    c.drawRightString(width - margin, 35, f"{number} / {total}")
    c.showPage()


def make_pdf(path: Path, pages: list[dict], family: str, *, scan: bool = False, seed: int = 0) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=(612, 792), invariant=1, pageCompression=1)
    c.setTitle("Document")
    c.setAuthor("Synthetic example")
    for i, page in enumerate(pages, 1):
        if page.get("blank"):
            c.showPage()
        else:
            draw_page(c, page, family, i, len(pages))
    c.save()
    raw = out.getvalue()
    if scan:
        import pypdfium2 as pdfium
        src = pdfium.PdfDocument(raw)
        scanned = io.BytesIO()
        c = canvas.Canvas(scanned, pagesize=(612, 792), invariant=1, pageCompression=1)
        c.setTitle("Document")
        c.setAuthor("Synthetic example")
        from reportlab.lib.utils import ImageReader
        for i, page in enumerate(src):
            image = page.render(scale=1.55).to_pil().convert("L")
            angle = random.Random(seed + i).uniform(-.6, .6)
            image = image.rotate(angle, resample=Image.Resampling.BICUBIC, fillcolor=245)
            image = image.filter(ImageFilter.GaussianBlur(.22))
            c.drawImage(ImageReader(image), 0, 0, width=612, height=792)
            c.showPage()
        c.save()
        src.close()
        raw = scanned.getvalue()
    path.write_bytes(raw)


def make_docx(path: Path, pages: list[dict]) -> None:
    doc = Document()
    props = doc.core_properties
    props.title = "Document"
    props.author = "Synthetic example"
    props.created = props.modified = datetime(2026, 1, 1, tzinfo=UTC)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.78)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Overused Grotesk", Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ["Title", "Subtitle", "Heading 1"]:
        doc.styles[name].font.name = "Overused Grotesk"
        doc.styles[name].font.color.rgb = RGBColor.from_string("000000")
    doc.styles["Title"].font.size = Pt(29)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    for index, page in enumerate(pages):
        if index:
            doc.add_page_break()
        p = doc.add_paragraph(page["organization"].upper())
        p.runs[0].font.color.rgb = RGBColor.from_string("3E18F9")
        p.runs[0].font.size = Pt(9)
        if page["title"]:
            doc.add_paragraph(page["title"], "Title")
        doc.add_paragraph(page["subtitle"], "Subtitle")
        for key, val in page.get("fields", []):
            p = doc.add_paragraph()
            p.add_run(key + ": ").bold = True
            p.add_run(val)
        for heading, text in page.get("sections", []):
            doc.add_paragraph(heading, "Heading 1")
            doc.add_paragraph(text)
        if table := page.get("table"):
            t = doc.add_table(rows=0, cols=len(table[0]))
            t.style = "Light Shading Accent 1"
            for row in table:
                for cell, val in zip(t.add_row().cells, row, strict=True):
                    cell.text = val
    foot = section.footer.paragraphs[0]
    foot.text = "SYNTHETIC EXAMPLE / NO REAL PERSON OR TRANSACTION"
    foot.runs[0].font.size = Pt(7)
    # Remove the stock Word template's title rule and theme font overrides.
    for element in doc.styles.element.iter():
        if element.tag == qn("w:pBdr"):
            element.getparent().remove(element)
        elif element.tag == qn("w:rFonts"):
            for key in list(element.attrib):
                if key.endswith("Theme"):
                    del element.attrib[key]
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                element.set(qn("w:" + attr), "Overused Grotesk")
    for paragraph_ in doc.paragraphs:
        for run in paragraph_.runs:
            run.font.name = "Overused Grotesk"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    neutral_zip(path)


def classify_records(output: Path) -> list[dict]:
    records = []
    for label_index, label in enumerate(CLASS_LABELS):
        for item in range(10):
            index = label_index * 10 + item + 1
            partition = "dev" if item < 2 else "test"
            family = f"classify-{partition}-layout-{item % 2 if partition == 'dev' else item % 4}"
            fmt = "docx" if item == 8 else "pdf"
            scan = item in (3, 7)
            count = 2 if item in (1, 4, 6, 9) else 1
            name = f"c{index:03d}"
            relative = f"datasets/files/{name}.{fmt}"
            pages = content(label, index, pages=count)
            if fmt == "docx":
                make_docx(output / relative, pages)
                family = "classify-test-office-1"
            else:
                make_pdf(output / relative, pages, family, scan=scan, seed=index)
            records.append({"id": name, "path": relative, "split": partition, "category": label, "page_count": count, "source_ids": [f"class-source-{index:03d}"], "template_family": family, "format": fmt, "scan": scan, "sha256": digest(output / relative)})
    return records


def packet_record(output: Path, index: int, *, demo: bool = False) -> dict:
    partition = "demo" if demo else ("dev" if index < 6 else "test")
    family = f"split-{partition}-layout-{index % (2 if partition == 'dev' else 4)}"
    name = "claim-packet" if demo else f"s{index+1:03d}"
    relative = "examples/split/claim-packet.pdf" if demo else f"datasets/files/{name}.pdf"
    specs = [("claim_form", 2), ("incident_report", 2 if index % 2 == 0 else 1), ("repair_estimate", 2), ("invoice", 2), ("invoice", 1), ("correspondence", 1)]
    if not demo and index % 3 == 0:
        specs.append(("other", 1))
    if not demo and index % 4 == 0:
        specs.insert(2, ("other", 1))
    writer = PdfWriter()
    segments = []
    sources = []
    pages_before = 0
    scan = not demo and index % 5 == 0
    with tempfile.TemporaryDirectory() as temp:
        for j, (label, length) in enumerate(specs):
            source = f"{partition}-packet-{index:03d}-source-{j:02d}"
            sources.append(source)
            # A distinct source index makes adjacent invoices visibly independent.
            pages = content(label, 1000 + index * 13 + j * 3, pages=length, domain="splitting")
            blank = label == "other" and j == 2
            if blank:
                pages = [{"blank": True}]
            temp_pdf = Path(temp) / f"{j}.pdf"
            make_pdf(temp_pdf, pages, family, scan=scan and not blank, seed=1000+index+j)
            for page in PdfReader(temp_pdf).pages:
                writer.add_page(page)
            numbers = list(range(pages_before + 1, pages_before + length + 1))
            segments.append({"category": label, "pages": numbers})
            pages_before += length
    path = output / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    writer.add_metadata({"/Title": "Document", "/Author": "Synthetic example"})
    writer.write(path)
    return {"id": name, "path": relative, "split": partition, "segments": segments, "page_count": pages_before, "source_ids": sources, "template_family": family, "format": "pdf", "scan": scan, "sha256": digest(path)}


def demos(output: Path) -> dict:
    records = []
    labels = CLASS_LABELS
    titles = ["Vendor invoice", "Scanned supply order", "Service agreement", "Candidate resume", "Investor presentation", "Observatory field guide"]
    for i, (label, title) in enumerate(zip(labels, titles, strict=True)):
        fmt = "docx" if label in ("contract", "resume") else "pptx" if label == "pitch_deck" else "pdf"
        relative = f"examples/classify/inbox/d{i+1:02d}.{fmt}"
        path = output / relative
        pages = content(label, 2100 + i, pages=2 if label in ("invoice", "contract") else 1)
        if fmt == "docx":
            make_docx(path, pages)
        elif fmt == "pdf":
            make_pdf(path, pages, f"demo-layout-{i}", scan=label == "purchase_order", seed=SEED)
        else:
            existing = ROOT / relative
            if not existing.exists():
                continue
            if path != existing:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(existing.read_bytes())
        records.append({"id": f"d{i+1:02d}", "path": relative, "title": title, "category": label, "page_count": 3 if fmt == "pptx" else len(pages), "source_ids": [f"demo-class-source-{i}"], "format": fmt, "scan": label == "purchase_order", "sha256": digest(path)})
    packet = packet_record(output, 99, demo=True)
    packet["title"] = "Property claim packet"
    return {"classify": records, "split": [packet]}


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2) + "\n")


def parity_fixtures(output: Path) -> None:
    """Related semantic equivalents are fixtures, never independent test samples."""
    deck = ROOT / "examples/classify/inbox/d05.pptx"
    if not deck.exists():
        return
    base = {"organization": "North Fern (fictional)", "reference": "PARITY-1", "date": "September 2026", "fields": []}
    pages = [
        {**base, "title": "Equipment maintenance for independent teams", "subtitle": "Seed investor presentation", "sections": [("September 2026", "North Fern is a fictional company. All figures are fictional.")]},
        {**base, "title": "A shared maintenance record", "subtitle": "North Fern", "sections": [("Customer problem", "Service teams lose time reconstructing equipment history from scattered spreadsheets."), ("Product", "North Fern combines service logs and maintenance schedules in one subscription product."), ("Illustrative annual price", "$12,000 per team.")]},
        {**base, "title": "Seed financing plan", "subtitle": "North Fern", "sections": [("Illustrative capital request", "$2 million. Fund twelve months of product development and a small paid pilot with service firms managing 50 to 500 assets."), ("Founding team", "Field operations and software engineering.")]},
    ]
    root = output / "examples/fixtures"
    root.mkdir(parents=True, exist_ok=True)
    make_pdf(root / "f01.pdf", pages, "parity-layout-2")
    make_docx(root / "f01.docx", pages)
    (root / "f01.pptx").write_bytes(deck.read_bytes())
    records = [{"path": f"examples/fixtures/f01.{fmt}", "format": fmt, "page_count": 3, "sha256": digest(root / f"f01.{fmt}"), "related_group": "north-fern-presentation", "category": "pitch_deck", "excluded_from_accuracy": True} for fmt in ("pdf", "docx", "pptx")]
    write_json(root / "manifest.json", records)


def build(output: Path) -> dict:
    fonts()
    classification = classify_records(output)
    split = [packet_record(output, i) for i in range(24)]
    demo = demos(output)
    parity_fixtures(output)
    write_json(output / "datasets/manifests/classify.json", classification)
    write_json(output / "datasets/manifests/split.json", split)
    write_json(output / "examples/demo-manifest.json", demo)
    return {"classify": classification, "split": split, "demo": demo}


def verify() -> None:
    expected = {name: json.loads((ROOT / f"datasets/manifests/{name}.json").read_text()) for name in ("classify", "split")}
    for name, rows in expected.items():
        ids = set()
        source_sets = {"dev": set(), "test": set()}
        family_sets = {"dev": set(), "test": set()}
        for row in rows:
            assert row["id"] not in ids
            ids.add(row["id"])
            path = ROOT / row["path"]
            assert digest(path) == row["sha256"], path
            current = source_sets[row["split"]]
            assert not current.intersection(row["source_ids"])
            current.update(row["source_ids"])
            family_sets[row["split"]].add(row["template_family"])
            if path.suffix == ".pdf":
                assert len(PdfReader(path).pages) == row["page_count"]
            if name == "split":
                assert [p for s in row["segments"] for p in s["pages"]] == list(range(1, row["page_count"] + 1))
        assert source_sets["dev"].isdisjoint(source_sets["test"])
        assert family_sets["dev"].isdisjoint(family_sets["test"])
    assert Counter(row["category"] for row in expected["classify"]) == {label: 10 for label in CLASS_LABELS}
    with tempfile.TemporaryDirectory() as temp:
        rebuilt = build(Path(temp))
        for name in ("classify", "split"):
            assert rebuilt[name] == expected[name], f"{name} failed byte-for-byte regeneration"
        assert rebuilt["demo"] == json.loads((ROOT / "examples/demo-manifest.json").read_text())
        assert (Path(temp) / "examples/fixtures/manifest.json").read_bytes() == (ROOT / "examples/fixtures/manifest.json").read_bytes()
    print("Verified hashes, source/template isolation, coverage, balance, and deterministic regeneration.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    if args.verify:
        verify()
    else:
        result = build(ROOT)
        print(f"Generated {len(result['classify'])} classification documents and {len(result['split'])} splitting packets.")
