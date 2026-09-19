"""Document invariants: source identity, exact pagination, and safe failures."""

from __future__ import annotations

import json
import os
import subprocess
import time
from pathlib import Path

import pytest
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from jev_docs.cache import cache_key, load_document
from jev_docs.conversion import find_soffice, render_page_png
from jev_docs.documents import parse_document, validate_pages
from jev_docs.errors import DocumentError
from jev_docs.ocr.base import OCRPage, OCRResult
from jev_docs.schemas import ParserInfo


def native_pdf(path: Path, *, blank: bool = False) -> Path:
    writer = canvas.Canvas(str(path), invariant=True)
    writer.drawString(72, 720, "INVOICE 1024. Total due 250 dollars.")
    writer.showPage()
    if blank:
        writer.showPage()
    writer.save()
    return path


def result(numbers: list[int], *, total: int = 2, texts: list[str] | None = None) -> OCRResult:
    return OCRResult(pages=[OCRPage(n, texts[i] if texts else "readable")
                            for i, n in enumerate(numbers)], total_pages=total,
                     parser=ParserInfo(name="fixture", version="1"))


@pytest.mark.parametrize("numbers,total", [([1], 2), ([2, 1], 2), ([1, 1], 2), ([1, 3], 2), ([1, 2], 3)])
def test_incomplete_or_reordered_pages_fail(tmp_path, numbers, total):
    with pytest.raises(DocumentError, match="pages|page"):
        validate_pages(result(numbers, total=total), tmp_path / "unused.pdf", 2)


def test_native_blank_preserved_cache_and_source_rename(tmp_path):
    source = native_pdf(tmp_path / "source.pdf", blank=True)
    first = parse_document(source, cache_dir=tmp_path / "cache")
    assert first.page_count == 2
    assert first.pages[0].number == 1 and "1024" in first.pages[0].text
    assert first.pages[1].number == 2 and first.pages[1].blank
    assert not first.metrics.ocr_cache_hit
    assert Path(first.canonical_path).read_bytes() == source.read_bytes()
    moved = tmp_path / "renamed.pdf"
    moved.write_bytes(source.read_bytes())
    second = parse_document(moved, cache_dir=tmp_path / "cache")
    assert second.metrics.ocr_cache_hit
    assert second.source_name == "renamed.pdf"
    assert second.metrics.ocr_ms == 0
    assert second.metrics.ocr_cost_usd == 0
    assert second.metrics.requests == []


def test_visible_page_without_text_is_not_blank(tmp_path):
    source = native_pdf(tmp_path / "source.pdf")
    with pytest.raises(DocumentError, match="visible content"):
        validate_pages(result([1], total=1, texts=[""]), source, 1)


def test_cache_invalidates_for_options_and_corrupt_records(tmp_path):
    source = native_pdf(tmp_path / "source.pdf")
    root = tmp_path / "cache"
    original = parse_document(source, cache_dir=root)
    assert cache_key({"tier": "agentic"}) != cache_key({"tier": "agentic_plus"})
    assert cache_key({"version": "1"}) != cache_key({"version": "2"})
    cache_record = next((root / "ocr").glob("*.json"))
    cache_record.write_text("{invalid")
    reparsed = parse_document(source, cache_dir=root)
    assert not reparsed.metrics.ocr_cache_hit
    assert reparsed.pages == original.pages
    record = json.loads(cache_record.read_text())
    record["pages"] = []
    cache_record.write_text(json.dumps(record))
    assert not parse_document(source, cache_dir=root).metrics.ocr_cache_hit


def test_cache_can_be_bypassed(tmp_path):
    source = native_pdf(tmp_path / "source.pdf")
    parse_document(source, cache_dir=tmp_path / "cache")
    assert not parse_document(source, cache_dir=tmp_path / "cache", use_cache=False).metrics.ocr_cache_hit


def test_mutable_alias_cache_expires(tmp_path):
    source = native_pdf(tmp_path / "source.pdf")
    root = tmp_path / "cache"
    document = parse_document(source, cache_dir=root)
    record = next((root / "ocr").glob("*.json"))
    old = time.time() - 25 * 3600
    os.utime(record, (old, old))
    canonical = Path(document.canonical_path)
    assert load_document(root, record.stem, canonical, max_age_seconds=24 * 3600) is None
    assert load_document(root, record.stem, canonical) is not None


def test_bad_empty_encrypted_and_unsupported_inputs(tmp_path):
    with pytest.raises(DocumentError, match="does not exist"):
        parse_document(tmp_path / "missing.pdf")
    source = tmp_path / "empty.pdf"
    source.touch()
    with pytest.raises(DocumentError, match="empty"):
        parse_document(source)
    source.write_bytes(b"not a PDF")
    with pytest.raises(DocumentError, match="could not be read"):
        parse_document(source, cache_dir=tmp_path / "cache")
    text = tmp_path / "notes.txt"
    text.write_text("invoice")
    with pytest.raises(DocumentError, match="Unsupported"):
        parse_document(text, cache_dir=tmp_path / "cache")
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    writer.write(source)
    with pytest.raises(DocumentError, match="Encrypted"):
        parse_document(source, cache_dir=tmp_path / "cache")


def test_render_preview_preserves_page_identity(tmp_path):
    source = native_pdf(tmp_path / "source.pdf", blank=True)
    preview = render_page_png(source, 2, tmp_path / "preview.png", scale=1)
    assert preview.is_file()
    with pytest.raises(DocumentError, match="outside"):
        render_page_png(source, 3, tmp_path / "invalid.png")


def test_office_failure_is_actionable_and_uses_isolated_profile(tmp_path, monkeypatch):
    from jev_docs import conversion

    source = tmp_path / "unsafe name.docx"
    source.write_bytes(b"invalid document")
    monkeypatch.setattr(conversion, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(conversion, "converter_version", lambda _: "fixture-office")
    monkeypatch.setattr(conversion, "font_fingerprint", lambda: {})
    commands = []

    def fail(command, **kwargs):
        commands.append(command)
        assert kwargs["timeout"] == 180
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        return subprocess.CompletedProcess(command, 1, stdout=b"", stderr=b"private content")

    monkeypatch.setattr(conversion.subprocess, "run", fail)
    with pytest.raises(DocumentError, match="could not convert") as error:
        parse_document(source, cache_dir=tmp_path / "cache")
    assert "private content" not in str(error.value)
    assert len(commands) == 1


@pytest.mark.skipif(not find_soffice(), reason="LibreOffice is an optional Office prerequisite")
def test_docx_and_pptx_canonical_pages(tmp_path):
    from docx import Document
    from pptx import Presentation

    word = Document()
    word.add_heading("Contract 4102", 0)
    word.add_paragraph("Services agreement, first page.")
    word.add_page_break()
    word.add_paragraph("Services agreement, second page.")
    word.save(tmp_path / "source.docx")
    deck = Presentation()
    for title in ("Quarterly Update", "Operating Plan"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Review costs and delivery schedule."
    deck.save(tmp_path / "source.pptx")
    for suffix in ("docx", "pptx"):
        parsed = parse_document(tmp_path / f"source.{suffix}", cache_dir=tmp_path / "cache")
        assert parsed.page_count == 2
        assert [page.number for page in parsed.pages] == [1, 2]
        assert Path(parsed.canonical_path).suffix == ".pdf"
        assert parsed.parser.conversion["profile"] == "isolated-temporary"
        again = parse_document(tmp_path / f"source.{suffix}", cache_dir=tmp_path / "cache")
        assert again.metrics.ocr_cache_hit
